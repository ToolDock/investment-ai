# -*- coding: utf-8 -*-
"""人手のコーディング（coding_sheet.xlsx）と、LLMの判定の一致度を出す。ローカル専用。
入力: 記入済みの coding_sheet.xlsx, coding_key.csv, llm_utterance_codes.csv, llm_response_scores.csv
使い方: python compare_coding.py [記入済みのxlsx=coding_sheet.xlsx]
出力: 標準出力（コードの一致率・Cohenのκ・混同行列、観点ごとのスピアマン順位相関と平均差）
"""
import csv
import sys
from collections import Counter

from openpyxl import load_workbook

from coding_defs import CODES, RUBRIC


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return po, (po - pe) / (1 - pe) if pe < 1 else float("nan")


def rank(v):
    order = sorted(range(len(v)), key=lambda i: v[i])
    r, i = [0.0] * len(v), 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and v[order[j + 1]] == v[order[i]]:
            j += 1
        for k in range(i, j + 1):
            r[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return r


def spearman(x, y):
    rx, ry = rank(x), rank(y)
    mx, my = sum(rx) / len(rx), sum(ry) / len(ry)
    num = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    den = (sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry)) ** 0.5
    return num / den if den else float("nan")


def num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def main():
    key = {k["id"]: (k["pid"], k["month"], k["turn"]) for k in read("coding_key.csv")}
    ws = load_workbook(sys.argv[1] if len(sys.argv) > 1 else "coding_sheet.xlsx", data_only=True)["コーディング"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    human = {key[r[0]]: r for r in rows if r[0] in key}
    llm_c = {(r["pid"], r["month"], r["turn"]): r for r in read("llm_utterance_codes.csv")}
    llm_s = {(r["pid"], r["month"], r["turn"]): r for r in read("llm_response_scores.csv")}

    pairs = [(h[5], llm_c[k]["code"]) for k, h in human.items() if h[5] and k in llm_c and llm_c[k]["code"]]
    print(f"== 発話のコード（{len(pairs)}件）")
    if pairs:
        a, b = zip(*pairs)
        po, k = kappa(a, b)
        print(f"一致率 {po:.2f} / Cohenのκ {k:.2f}")
        conf = Counter(pairs)
        print("人手→LLM の食い違い（多い順）:")
        for (x, y), c in conf.most_common():
            if x != y:
                print(f"  {x}→{y}: {c}")
        low = [(k_, llm_c[k_]["confidence"]) for k_ in human if k_ in llm_c]
        print("LLMの確信度の分布:", dict(Counter(c for _, c in low)))

    print("\n== 応答の評価（観点ごと。順位相関と、平均の差=LLM−人手）")
    for i, (key_, name, _) in enumerate(RUBRIC):
        x, y = [], []
        for k, h in human.items():
            hv, lv = num(h[6 + i]), num(llm_s.get(k, {}).get(key_))
            if hv is not None and lv is not None:
                x.append(hv)
                y.append(lv)
        if len(x) >= 5:
            print(f"{name}: n={len(x)} ρ={spearman(x, y):.2f} 平均差={sum(y) / len(y) - sum(x) / len(x):+.2f}")
        else:
            print(f"{name}: 比較できる件数が少ない（n={len(x)}）")


if __name__ == "__main__":
    main()
